import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from typing import List, Dict, Any

class ExcelExporter:
    """결과 리포트 Excel 생성기"""

    COLUMN_MAPPING = {
        "title": "공고명",
        "competent_authority": "소관부처",
        "local_government": "지자체",
        "application_period": "신청기간",
        "project_overview": "사업개요",
        "budget": "사업비",
        "application_method": "신청방법",
        "contact": "문의처",
        "bid_strategy": "수주전략",
        "url": "원문 링크",
    }

    @staticmethod
    def export(data_list: List[Dict[str, Any]], output_path: str):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        if not data_list:
            print("[Info] 내보낼 데이터가 없습니다.")
            return

        rows = sorted(
            data_list,
            key=lambda row: row.get("relevance_score", 0) or 0,
            reverse=True,
        )

        available_cols = [col for col in ExcelExporter.COLUMN_MAPPING.keys() if any(col in row for row in rows)]
        headers = ["번호"] + [ExcelExporter.COLUMN_MAPPING[col] for col in available_cols]

        wb = Workbook()
        ws = wb.active
        ws.title = "AI·로봇 공고 목록"
        ws.append(headers)

        for idx, row in enumerate(rows, start=1):
            ws.append([idx] + [row.get(col) if row.get(col) not in (None, "") else "확인필요" for col in available_cols])

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="맑은 고딕", size=9)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = 0
            for cell in col:
                cell.font = body_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('euc-kr', errors='ignore')))

            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        wb.save(output_path)
        print(f"[Success] 엑셀 리포트 생성 완료: {output_path}")
